# -*- coding: utf-8 -*-
"""阶段4: 用 LibreOffice 把老格式 .ppt 逐个转 .pptx 再提取(textutil/python-pptx 都读不了老 .ppt)。
逐个转换、随转随取随清,一个损坏不带累整批,同名文件不互相污染。
用法: python3 convert_legacy.py --src <素材包> --out <库> [--soffice <路径>]
若 LibreOffice 未装: brew install --cask libreoffice"""
import os, sys, argparse, subprocess, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import clean, judge_quality, load_manifest, save_manifest, rec_key, flat_name, find_soffice

def extract_pptx(p):
    from pptx import Presentation
    prs = Presentation(p); parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame: parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append("\t".join(c.text for c in row.cells))
    n = len(prs.slides); full = clean("\n".join(parts)); chars = len(full.strip())
    return full, n, chars, judge_quality(chars, n, "pptx")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--soffice", default=None, help="LibreOffice 路径,留空自动检测")
    a = ap.parse_args()
    soff = a.soffice or find_soffice()
    assert soff and os.path.exists(soff), (
        "LibreOffice 未找到。\n"
        "  macOS: brew install --cask libreoffice\n"
        "  Linux: apt install libreoffice (或 dnf/pacman)\n"
        "  Windows: 从 libreoffice.org 下载安装"
    )
    RAW = os.path.join(a.out, "_raw", "all"); os.makedirs(RAW, exist_ok=True)
    manifest = load_manifest(a.out)
    targets = [r for r in manifest if r.get("quality") == "needs_conversion"]
    if not targets:
        print("无待转换 .ppt"); return
    print(f"待转换 {len(targets)} 个, LibreOffice 逐个转换中(首次启动慢)...")
    TMP = tempfile.mkdtemp(prefix="corpus_ppt_conv_")
    ok = 0
    try:
        for rec in targets:
            s = os.path.join(a.src, rec_key(rec))
            if not os.path.exists(s):
                print(f"[SKIP] 原始文件不在: {rec['name'][:40]}"); continue
            subprocess.run([soff, "--headless", "--convert-to", "pptx", "--outdir", TMP, s],
                           capture_output=True, timeout=300)
            base = os.path.splitext(os.path.basename(rec["name"]))[0]
            pptx = os.path.join(TMP, base + ".pptx")
            if not os.path.exists(pptx):
                rec["quality"] = "failed"; rec["error"] = "convert failed(源文件可能损坏)"
                print(f"[FAIL] {rec['name'][:40]}"); continue
            try:
                full, n, chars, q = extract_pptx(pptx)
            except Exception as e:
                rec["quality"] = "failed"; rec["error"] = str(e)
                os.remove(pptx); continue
            os.remove(pptx)   # 随转随清,防同名 base 互相污染
            rp = os.path.join(RAW, flat_name(rec_key(rec)) + ".txt")
            open(rp, "w", encoding="utf-8").write(full)
            rec.update({"unit": "slides", "count": n, "chars": chars, "quality": q,
                        "raw": os.path.relpath(rp, a.out), "converted": True})
            rec.pop("error", None)
            ok += 1
            print(f"[OK] {q:6} {n}页 {chars}字  {rec['name'][:40]}")
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    save_manifest(a.out, manifest)
    print(f"\n转换成功 {ok}/{len(targets)}。下一步: build.py 重建库")

if __name__ == "__main__":
    main()
