# -*- coding: utf-8 -*-
"""阶段1+2: 扫描类型分布 + 分层提取文本到 _raw/all/, 增量合并进 manifest.json。
用法: python3 extract.py --src <素材包目录> --out <输出库目录> [--force]
零成本(无 LLM)。老格式 .ppt 标 needs_conversion 留给 convert_legacy.py。

增量语义(重要):
- 默认可安全重跑。已妥善处理的记录(含 OCR 转正/乱码修复/格式转换的成果)原样保留,
  只提取新文件和上次 failed/unsupported 的文件。往素材包丢新文件后重跑本脚本即完成增量入库。
- --force 对 src 里现存的所有文件强制重提(会重置这些文件的 OCR/修复状态,慎用)。
- 递归扫描子目录,以相对路径为唯一键,同名不同目录/不同扩展名互不覆盖。"""
import os, sys, argparse, zipfile, subprocess, re, tempfile, shutil, html
from collections import Counter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (clean, judge_quality, classify, load_manifest, save_manifest,
                    manifest_index, rec_key, is_settled, raw_filename, safe_join,
                    find_soffice, secure_makedirs, secure_write_text, RETRYABLE)

def extract_pdf(p):
    import fitz
    doc = fitz.open(p); n = doc.page_count
    parts = [pg.get_text("text") for pg in doc]
    doc.close()
    full = clean("\n".join(parts)); chars = len(full.strip())
    return full, {"unit": "pages", "count": n, "chars": chars, "quality": judge_quality(chars, n, "pdf")}

def extract_pptx(p):
    from pptx import Presentation
    prs = Presentation(p); parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame: parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append("\t".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            parts.append("[备注] " + slide.notes_slide.notes_text_frame.text)
    n = len(prs.slides); full = clean("\n".join(parts)); chars = len(full.strip())
    return full, {"unit": "slides", "count": n, "chars": chars, "quality": judge_quality(chars, n, "pptx")}

def extract_docx(p):
    import docx
    d = docx.Document(p); parts = [par.text for par in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    full = clean("\n".join(parts)); chars = len(full.strip())
    return full, {"unit": "paragraphs", "count": len(d.paragraphs), "chars": chars,
                  "quality": "text" if chars >= 200 else "sparse"}

def extract_doc(p):
    """处理 .doc 老格式 Word。macOS 用 textutil(快),其他平台 fallback 到 LibreOffice。"""
    if sys.platform == "darwin" and shutil.which("textutil"):
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", p],
                           capture_output=True, timeout=60)
        full = clean(r.stdout.decode("utf-8", "ignore"))
    else:
        soff = find_soffice()
        if not soff:
            raise RuntimeError(".doc 需要 textutil(macOS自带) 或 LibreOffice(其他平台)")
        tmp = tempfile.mkdtemp()
        try:
            subprocess.run([soff, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, p],
                           capture_output=True, timeout=120)
            txt = os.path.join(tmp, os.path.splitext(os.path.basename(p))[0] + ".txt")
            full = clean(open(txt, encoding="utf-8").read()) if os.path.exists(txt) else ""
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    chars = len(full.strip())
    return full, {"unit": "doc", "count": 1, "chars": chars,
                  "quality": "text" if chars >= 200 else "sparse"}

EPUB_CHAPTER_CAP = 50 * 1024 * 1024   # 单章节解压后上限 50MB,防 zip bomb 耗尽内存

def extract_epub(p):
    z = zipfile.ZipFile(p)
    htmls = sorted(n for n in z.namelist() if n.lower().endswith((".html", ".xhtml", ".htm")))
    parts = []
    for h in htmls:
        if z.getinfo(h).file_size > EPUB_CHAPTER_CAP:
            raise ValueError(f"epub 章节解压后超 50MB,疑似 zip bomb: {h}")
        raw = z.read(h).decode("utf-8", "ignore")
        raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
        raw = html.unescape(raw)  # 先解码全部 HTML 实体(含 &#x4E2D; 数字字符引用),再剥标签
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        parts.append(raw)
    full = re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]{2,}", " ", clean("\n".join(parts))))
    return full, {"unit": "chapters", "count": len(htmls), "chars": len(full.strip()), "quality": "text"}

EXTRACTORS = {".pdf": extract_pdf, ".pptx": extract_pptx, ".docx": extract_docx,
              ".doc": extract_doc, ".epub": extract_epub}

def walk_src(src, out):
    """递归列出素材文件的相对路径。跳过隐藏文件/目录,跳过嵌在 src 里的输出库目录。"""
    out_abs = os.path.abspath(out)
    rels = []
    for root, dirs, files in os.walk(src):
        dirs[:] = [d for d in dirs if not d.startswith(".")
                   and os.path.abspath(os.path.join(root, d)) != out_abs]
        for f in files:
            if f.startswith("."):
                continue
            rels.append(os.path.relpath(os.path.join(root, f), src).replace(os.sep, "/"))
    return sorted(rels)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--force", action="store_true",
                    help="对 src 现存文件强制重提(会重置其 OCR/修复状态)")
    a = ap.parse_args()
    try:
        RAW = safe_join(a.out, os.path.join("_raw", "all"))
    except ValueError as e:
        sys.exit(f"拒绝不安全的输出路径: {e}")
    secure_makedirs(a.out, os.path.join("_raw", "all"))

    rels = walk_src(a.src, a.out)
    by_ext = Counter(os.path.splitext(r)[1].lower() for r in rels)
    print("=== 类型分布 ===")
    for ext, n in by_ext.most_common():
        print(f"  {ext or '(无)'}: {n}")

    manifest = load_manifest(a.out)
    idx = manifest_index(manifest)
    # epub/mobi 去重: 同路径同名 epub 在则 mobi skip
    epub_bases = {os.path.splitext(r)[0] for r in rels if r.lower().endswith(".epub")}

    raw_counts = Counter(r.get("raw") for r in manifest if r.get("raw"))
    new = kept = retried = failed = 0
    for i, rel in enumerate(rels, 1):
        name = os.path.basename(rel)
        ext = os.path.splitext(name)[1].lower()
        p = os.path.join(a.src, rel)
        old = idx.get(rel)
        if old is None and name in idx:
            # 兼容 v0.1 老键: 同名记录须 ext+size 一致才绑定,防同名不同文件错绑状态
            cand = idx[name]
            try:
                same_file = cand.get("ext") == ext and cand.get("size") == os.path.getsize(p)
            except OSError:
                same_file = False
            if same_file:
                old = cand
        desired_raw = os.path.join("_raw", "all", raw_filename(rel))
        refresh_raw = bool(old and old.get("raw") and (
            raw_counts[old["raw"]] > 1
            or os.path.normpath(old["raw"]) != os.path.normpath(desired_raw)
        ))
        if old is not None and not a.force and not refresh_raw and is_settled(old, a.out):
            old.setdefault("rel", rel)
            kept += 1
            continue
        if old is not None and (old.get("quality") in RETRYABLE or refresh_raw):
            retried += 1

        rec = old if old is not None else {}
        rec.update({"name": name, "rel": rel, "ext": ext})
        rec.setdefault("category", classify(name))
        try:
            rec["size"] = os.path.getsize(p)
            if ext == ".ppt":
                rec.update({"quality": "needs_conversion", "note": "老格式PPT,需LibreOffice转换"})
                if refresh_raw:
                    rec.pop("raw", None); rec.pop("converted", None)
            elif ext == ".mobi":
                if os.path.splitext(rel)[0] in epub_bases:
                    rec.update({"quality": "skip_duplicate", "note": "与epub同书"})
                else:
                    rec.update({"quality": "unsupported", "note": "mobi无epub版,未提取"})
            elif ext in EXTRACTORS:
                full, meta = EXTRACTORS[ext](p)
                raw_rel = os.path.join("_raw", "all", raw_filename(rel))
                secure_write_text(a.out, raw_rel, full, create_parent=True)
                rec.update(meta); rec["raw"] = raw_rel.replace(os.sep, "/")
                rec.pop("error", None); rec.pop("ocr", None); rec.pop("garble_fixed", None)
            else:
                rec.update({"quality": "unsupported"})
        except Exception as e:
            rec.update({"quality": "failed", "error": f"{type(e).__name__}: {e}"})
            failed += 1
            print(f"[FAIL] {rel[:48]}: {e}")
        if old is None:
            manifest.append(rec); idx[rel] = rec; new += 1
        if i % 25 == 0:
            print(f"  ...{i}/{len(rels)}")

    save_manifest(a.out, manifest)
    qc = Counter(r.get("quality") for r in manifest)
    print("\n=== 提取完成 ===")
    print(f"本次: 新增 {new} | 保留已处理 {kept} | 重试 {retried} | 失败 {failed}")
    print("全库质量分布:", dict(qc))
    print("全库类目分布:", dict(Counter(r.get("category", "?") for r in manifest)))
    if new or retried:
        print("\n下一步: python3 scripts/build.py --out <库> 重建归档与索引")

if __name__ == "__main__":
    main()
